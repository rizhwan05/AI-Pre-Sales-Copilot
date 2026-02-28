import os
import io
import fitz  # PyMuPDF
from docx import Document
import openpyxl
from fastapi import HTTPException, UploadFile
from starlette.concurrency import run_in_threadpool


async def extract_text_from_upload(file: UploadFile) -> str:
    """
    Extract text from a FastAPI UploadFile object.
    Supports PDF, DOCX, XLSX, and plain text files.
    Returns the extracted text as a string.
    Raises HTTPException if extraction fails.
    """
    try:
        # Read file content as bytes
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="Empty file")

        # Detect file extension
        ext = os.path.splitext(file.filename)[1].lower()
        text = ""

        if ext == ".pdf":
            text = await run_in_threadpool(_extract_pdf_text, content)

        elif ext == ".docx":
            text = await run_in_threadpool(_extract_docx_text, content)

        elif ext in [".xls", ".xlsx"]:
            text = await run_in_threadpool(_extract_xlsx_text, content)

        else:
            # Treat as plain text
            text = content.decode("utf-8", errors="replace")

        if not text.strip():
            raise HTTPException(status_code=400, detail="No text extracted")

        return text

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Text extraction failed: {str(e)}")


def _extract_pdf_text(content: bytes) -> str:
    text = ""
    pdf = fitz.open(stream=content, filetype="pdf")
    for page in pdf:
        text += page.get_text()
    pdf.close()
    return text


def _extract_docx_text(content: bytes) -> str:
    text = ""
    doc = Document(io.BytesIO(content))
    for p in doc.paragraphs:
        text += p.text + "\n"
    return text


def _extract_xlsx_text(content: bytes) -> str:
    text = ""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell:
                    text += str(cell) + " "
            text += "\n"
    return text